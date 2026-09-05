"""Document readers for the Aali agent: PDF, Word, Excel.

These wrap pypdf / python-docx / openpyxl. They are pure-Python and safe:
they only read files inside the agent workspace (enforced by the caller in
file_tools.py, which resolves paths against the workspace root).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


def read_pdf(path: Path, *, max_pages: int = 40, max_chars: int = 60_000) -> dict[str, Any]:
    """Extract text from a PDF, page by page."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    total = len(reader.pages)
    chunks: list[str] = []
    used = 0
    for i, page in enumerate(reader.pages[:max_pages]):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - a corrupt page should not kill the read
            text = ""
        text = text.strip()
        if not text:
            continue
        header = f"— صفحة {i + 1} —\n"
        chunks.append(header + text)
        used += len(text) + len(header)
        if used >= max_chars:
            break
    body = "\n\n".join(chunks)[:max_chars]
    return {
        "pages_total": total,
        "pages_read": min(total, max_pages),
        "text": body,
        "truncated": used >= max_chars or total > max_pages,
    }


def read_docx(path: Path, *, max_chars: int = 60_000) -> dict[str, Any]:
    """Extract paragraphs and table text from a Word document."""
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                parts.append(line)
    body = "\n".join(parts)[:max_chars]
    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "text": body,
        "truncated": len(body) >= max_chars,
    }


def read_xlsx(path: Path, *, max_rows: int = 500, max_chars: int = 60_000) -> dict[str, Any]:
    """Extract every sheet as pipe-separated rows."""
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out: list[str] = []
    used = 0
    for sheet in workbook.worksheets:
        out.append(f"=== ورقة: {sheet.title} ===")
        used += 30
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= max_rows:
                out.append(f"(… تم اقتطاع {sheet.title} بعد {max_rows} صف)")
                break
            cells = ["" if v is None else str(v) for v in row]
            line = " | ".join(cells).rstrip(" |")
            if line:
                out.append(line)
                used += len(line) + 1
            if used >= max_chars:
                break
        if used >= max_chars:
            break
    workbook.close()
    body = "\n".join(out)[:max_chars]
    return {
        "sheets": len(workbook.sheetnames),
        "text": body,
        "truncated": used >= max_chars,
    }
