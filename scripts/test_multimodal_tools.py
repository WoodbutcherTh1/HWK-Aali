"""Functional test for the new multimodal tools: read_document + analyze_video.

Creates sample files in a temp workspace, runs the tools through
execute_tool, prints results. Exit code 0 = all good.
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "file-agent"))

from file_agent.file_tools import execute_tool  # noqa: E402


def make_sample_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("تقرير تجريبي", level=1)
    document.add_paragraph("هذه فقرة تجريبية باللغة العربية لاختبار القارئ.")
    document.add_paragraph("Second paragraph in English for mixed-direction testing.")
    document.save(str(path))


def make_sample_xlsx(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "المبيعات"
    sheet.append(["الشهر", "المبيعات"])
    sheet.append(["يناير", 1200])
    sheet.append(["فبراير", 1500])
    workbook.save(str(path))


def make_sample_video(path: Path, seconds: int = 4) -> None:
    import cv2
    import numpy as np

    writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), 10, (640, 360))
    for i in range(seconds * 10):
        frame = np.full((360, 640, 3), 245, dtype=np.uint8)
        # Alternate "slides" with big readable text (OCR-able)
        text = "HWK TEST FRAME" if (i // 10) % 2 == 0 else "AALI VIDEO OCR"
        cv2.putText(frame, text, (60, 190), cv2.FONT_HERSHEY_SIMPLEX, 1.6, (20, 20, 20), 3)
        writer.write(frame)
    writer.release()


def main() -> int:
    failures = 0
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)

        make_sample_docx(tmp_path / "report.docx")
        make_sample_xlsx(tmp_path / "sales.xlsx")
        make_sample_video(tmp_path / "clip.mp4")

        for name in ("report.docx", "sales.xlsx"):
            result = execute_tool("read_document", {"path": name}, tmp_path)
            ok = result.get("ok")
            body = result.get("result", {})
            text = body.get("text", "")
            print(f"[read_document {name}] ok={ok} chars={len(text)} truncated={body.get('truncated')}")
            print("   sample:", text[:110].replace("\n", " | "))
            if not ok or len(text) < 10:
                failures += 1

        # unsupported type must fail gracefully
        (tmp_path / "x.exe").write_bytes(b"MZ")
        result = execute_tool("read_document", {"path": "x.exe"}, tmp_path)
        print("[read_document x.exe] blocked =", not result["ok"])
        if result["ok"]:
            failures += 1

        result = execute_tool(
            "analyze_video", {"path": "clip.mp4", "frame_interval": 1.0, "max_frames": 4}, tmp_path
        )
        ok = result.get("ok")
        body = result.get("result", {})
        frames = body.get("frames_ocr", [])
        print(f"[analyze_video clip.mp4] ok={ok} sampled={body.get('frames_sampled')} ocr_hits={len(frames)}")
        for hit in frames[:2]:
            print("   frame@", hit.get("at_sec"), "→", hit.get("text", "")[:60].replace("\n", " / "))
        if not ok or not frames:
            failures += 1

    print("FAILURES:", failures)
    return 0 if failures == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
